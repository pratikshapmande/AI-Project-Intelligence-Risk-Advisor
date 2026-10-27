from pathlib import Path
from io import BytesIO
import pandas as pd
from docx import Document
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".csv", ".txt"}

def extract_document(filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}. Use PDF, DOCX, CSV, or TXT.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    metadata = {}
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        parts = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(f"[Page {page_num}]\n{page_text}")
        text = "\n\n".join(parts)
        metadata = {"page_count": len(reader.pages)}
    elif suffix == ".docx":
        doc = Document(BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
        metadata = {"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)}
    elif suffix == ".csv":
        df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        text = df.to_csv(index=False)
        metadata = {"row_count": int(len(df)), "columns": [str(c) for c in df.columns]}
    else:
        text = content.decode("utf-8-sig", errors="replace")
    if not text.strip():
        raise ValueError("No readable text was extracted. Scanned/image-only PDFs may require OCR, which is not included in Milestone 1.")
    return {
        "source": Path(filename).name,
        "file_type": suffix[1:],
        "text": text.strip(),
        "metadata": metadata,
    }

def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    """Split text into character-based chunks with overlap, preferring a nearby word boundary."""
    if chunk_size < 1 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("chunk_size must be positive and overlap must be between 0 and chunk_size.")
    clean = " ".join(text.split())
    if not clean:
        return []
    chunks, start, length = [], 0, len(clean)
    while start < length:
        end = min(start + chunk_size, length)
        if end < length:
            boundary = clean.rfind(" ", start + int(chunk_size * 0.65), end)
            if boundary > start:
                end = boundary
        piece = clean[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= length:
            break
        start = max(end - overlap, start + 1)
    return chunks
