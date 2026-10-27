from io import BytesIO, StringIO
from pathlib import Path
import pandas as pd
from pypdf import PdfReader
from docx import Document


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from PDF, DOCX, CSV or TXT content."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        pages = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[Page {page_num}]\n{page_text}")
        return "\n\n".join(pages)

    if suffix == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs)

    if suffix == ".csv":
        # Try common encodings so CSV files exported from Excel are more likely to work.
        decoded = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("Could not decode this CSV file.")
        frame = pd.read_csv(StringIO(decoded), dtype=str, keep_default_na=False)
        return frame.to_csv(index=False)

    if suffix == ".txt":
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode this text file.")

    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")
